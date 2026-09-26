"""A patient's whole file as a Word document (.docx): to print, to send, or to keep outside the
system. Written in English with the data as entered (Arabic names stay Arabic)."""

import io

from django.utils import timezone, translation

from apps.billing.models import account
from apps.charting.forms import HISTORY_FIELDS
from apps.charting.models import ToothState

TABLE_STYLE = "Light Grid Accent 1"


def _text(value):
    if value is None or value == "":
        return ""
    if hasattr(value, "strftime"):
        if hasattr(value, "hour"):
            return timezone.localtime(value).strftime("%d/%m/%Y %I:%M %p") if timezone.is_aware(value) \
                else value.strftime("%d/%m/%Y %I:%M %p")
        return value.strftime("%d/%m/%Y")
    if isinstance(value, bool):
        return "Yes" if value else "No"
    return str(value)


def _field_value(obj, name):
    field = obj._meta.get_field(name)
    if field.many_to_many:
        return ", ".join(str(item) for item in getattr(obj, name).all())
    if field.choices:
        return str(getattr(obj, f"get_{name}_display")())
    return _text(getattr(obj, name))


def _table(document, headers, rows):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = TABLE_STYLE
    for cell, header in zip(table.rows[0].cells, headers):
        cell.text = str(header)
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = _text(value)
    document.add_paragraph()
    return table


def _facts(document, pairs):
    """Label / value lines, skipping empty values."""
    rows = [(label, value) for label, value in pairs if value not in (None, "", "—")]
    if rows:
        table = document.add_table(rows=0, cols=2)
        table.style = TABLE_STYLE
        for label, value in rows:
            cells = table.add_row().cells
            cells[0].text, cells[1].text = str(label), _text(value)
        document.add_paragraph()


def patient_docx(patient):
    from docx import Document
    from docx.shared import Pt

    with translation.override("en"):
        document = Document()
        document.styles["Normal"].font.name = "Calibri"
        document.styles["Normal"].font.size = Pt(10)
        document.add_heading(f"{patient.full_name} — {patient.file_number}", level=0)
        document.add_paragraph(f"Patient file exported on {timezone.localtime():%d/%m/%Y %I:%M %p}")

        document.add_heading("Personal data", level=1)
        names = ["file_number", "full_name", "id_type", "national_id", "birth_date", "gender", "marital_status",
                 "phone_primary", "phone_secondary", "address", "city", "governorate", "occupation", "referral_source",
                 "referral_details", "assigned_dentist", "status", "out_reason", "out_notes", "registered_on",
                 "missing_teeth", "missing_teeth_notes", "notes"]
        _facts(document, [(patient._meta.get_field(n).verbose_name.capitalize(), _field_value(patient, n))
                          for n in names if _has_field(patient, n)])

        exam = patient.examinations.prefetch_related("conditions").first()
        if exam is not None:
            document.add_heading("Medical and dental history", level=1)
            document.add_paragraph(f"From the examination of {_text(exam.exam_date)}")
            _facts(document, [(exam._meta.get_field(n).verbose_name.capitalize(), _field_value(exam, n))
                              for n in HISTORY_FIELDS if _has_field(exam, n)])

        states = list(patient.tooth_states.order_by("tooth"))
        if states:
            document.add_heading("Dental chart", level=1)
            rows = []
            for state in states:
                findings = [label for flag, label in ((state.caries, "caries"), (state.filled, "filled"),
                                                      (state.rct, "root canal"), (state.crown, "crown"),
                                                      (state.hopeless, "hopeless")) if flag]
                if state.status != ToothState.Status.PRESENT or findings:
                    rows.append((state.tooth, state.get_status_display(), ", ".join(findings)))
            _table(document, ["Tooth", "Status", "Findings"], rows)

        plans = list(patient.treatment_plans.select_related("dentist").prefetch_related("items__step_type"))
        if plans:
            document.add_heading("Treatment plans", level=1)
            for plan in plans:
                document.add_heading(f"{plan.title} ({plan.get_status_display()}, {_text(plan.created_at)})", level=2)
                for section in plan.sections():
                    document.add_paragraph(str(section["label"]), style="Intense Quote")
                    _table(document, ["Procedure", "Teeth", "Details", "Status"],
                           [(i.step_type, i.teeth, i.details, i.get_status_display()) for i in section["items"]])

        steps = list(patient.treatment_steps.select_related("step_type", "operator", "supervisor")
                     .order_by("performed_at"))
        if steps:
            document.add_heading("Treatment log", level=1)
            _table(document, ["Date", "Treatment", "Teeth", "Operator", "Supervisor", "Notes"],
                   [(s.performed_at, s.step_type, s.teeth, s.operator, s.supervisor, s.notes) for s in steps])

        surgeries = list(patient.surgeries.select_related("operator_1", "instructor")
                         .prefetch_related("sites__implant_system").order_by("date"))
        if surgeries:
            document.add_heading("Surgeries and implants", level=1)
            rows = []
            for surgery in surgeries:
                for site in surgery.sites.all():
                    rows.append((surgery.date, surgery.number, site.tooth, ", ".join(site.procedure_labels()),
                                 site.implant_label if site.has_implant else "", surgery.operator_1, surgery.instructor))
            _table(document, ["Date", "Surgery", "Tooth", "Procedures", "Implant", "Operator", "Supervisor"], rows)

        appointments = list(patient.appointments.select_related("room", "dentist").order_by("scheduled_at"))
        if appointments:
            document.add_heading("Appointments", level=1)
            _table(document, ["Date", "Status", "Room", "Dentist", "Purpose"],
                   [(a.scheduled_at, a.get_status_display(), a.room, a.dentist, a.what) for a in appointments])

        money = account(patient)
        if money["rows"] or money["payments"]:
            document.add_heading("Services and payments", level=1)
            _table(document, ["Date", "Service", "Price", "Discount", "Net", "Paid", "Left"],
                   [(r["charge"].charged_on, r["charge"].service, r["charge"].price, r["charge"].discount_amount,
                     r["charge"].net, r["paid"], r["left"]) for r in money["rows"]])
            _table(document, ["Paid on", "Amount", "Method", "Notes"],
                   [(p.paid_on, p.amount, p.get_method_display(), p.notes) for p in money["payments"]])
            _facts(document, [("Total", money["net"]), ("Paid", money["paid"]), ("Balance", money["balance"])])

        labs = list(patient.lab_requests.select_related("work_type", "lab", "dentist"))
        if labs:
            document.add_heading("Lab requests", level=1)
            _table(document, ["Date", "Work", "Teeth", "Lab", "Dentist", "Status"],
                   [(lab.created_at, lab.work_type, lab.teeth, lab.lab, lab.dentist, lab.get_status_display())
                    for lab in labs])

        output = io.BytesIO()
        document.save(output)
        output.seek(0)
        return output


def _has_field(obj, name):
    try:
        obj._meta.get_field(name)
    except Exception:
        return False
    return True
