import io
import shutil
import tempfile
from datetime import date, timedelta

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings

from apps.core.testing import PASSWORD, make_dentist, make_patient, make_user, setup_clinic
from apps.patients.models import Lead, Patient, PatientDocument, PatientRelation, ReferralSource
from apps.scheduling.models import Appointment
from django.utils import timezone

MEDIA = tempfile.mkdtemp()


@override_settings(MEDIA_ROOT=MEDIA)
class PatientRegistrationTests(TestCase):
    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        shutil.rmtree(MEDIA, ignore_errors=True)

    def setUp(self):
        self.branch = setup_clinic()
        self.secretary = make_user("sec", "secretary")
        self.dentist = make_dentist("dentist")
        self.client.login(username="sec", password=PASSWORD)
        self.facebook = ReferralSource.objects.get(name_en="Facebook")

    def form_data(self, **overrides):
        data = {
            "full_name": "محمد أحمد علي حسن",
            "id_type": "nid",
            "national_id": "29001150101234",
            "phone_primary": "01001234567",
            "preferred_phone": "primary",
            "missing_teeth": "single",
            "referral_source": self.facebook.pk,
            "assigned_dentist": self.dentist.pk,
            # Required at registration since round 13 (the birth date, gender and governorate come from the ID).
            "phone_secondary": "01101234567",
            "marital_status": "married",
            "occupation": "مهندس",
            "city": "مدينة نصر",
        }
        data.update(overrides)
        return data

    def test_register_patient_with_id_scan(self):
        data = self.form_data(id_front=SimpleUploadedFile("id.jpg", b"\xff\xd8\xff fake-image", content_type="image/jpeg"))
        response = self.client.post("/patients/new/", data)
        patient = Patient.objects.get()
        self.assertRedirects(response, f"/patients/{patient.pk}/", fetch_redirect_response=False)
        self.assertEqual(patient.file_number, f"CIA-{patient.pk:05d}")
        self.assertEqual(patient.birth_date, date(1990, 1, 15))  # read from the national ID
        self.assertEqual(patient.gender, "M")
        self.assertEqual(patient.assigned_dentist, self.dentist)
        self.assertEqual(patient.created_by, self.secretary)
        document = PatientDocument.objects.get()
        self.assertEqual(document.kind, PatientDocument.Kind.ID_FRONT)
        # The scan is served only through the protected view.
        self.assertEqual(self.client.get(document.file.url).status_code, 200)
        self.client.logout()
        self.assertEqual(self.client.get(document.file.url).status_code, 302)

    def test_duplicate_national_id_is_rejected(self):
        existing = make_patient(self.branch, nid="29001150101234", phone="01112223334")
        response = self.client.post("/patients/new/", self.form_data())
        self.assertEqual(Patient.objects.count(), 1)
        self.assertContains(response, existing.file_number)

    def test_duplicate_primary_phone_is_rejected_whatever_the_format(self):
        make_patient(self.branch, nid="28501010101235", phone="01001234567")
        for typed in ["+20 100 123 4567", "٠١٠٠١٢٣٤٥٦٧"]:
            response = self.client.post("/patients/new/", self.form_data(phone_primary=typed))
            self.assertEqual(Patient.objects.count(), 1, typed)
            self.assertIn("phone_primary", response.context["form"].errors)

    def test_invalid_national_id(self):
        response = self.client.post("/patients/new/", self.form_data(national_id="123"))
        self.assertIn("national_id", response.context["form"].errors)

    def test_referral_by_patient_choosing_the_patient_is_optional(self):
        # Round 13: "referred by someone we have" no longer asks to type the referring patient.
        by_patient = ReferralSource.objects.get(asks_for_patient=True)
        referrer = make_patient(self.branch, nid="28501010101235", phone="01112223334")
        self.client.post("/patients/new/", self.form_data(referral_source=by_patient.pk))
        patient = Patient.objects.get(national_id="29001150101234")
        self.assertIsNone(patient.referred_by)
        patient.delete()
        self.client.post("/patients/new/", self.form_data(referral_source=by_patient.pk, referred_by_lookup="01112223334"))
        self.assertEqual(Patient.objects.get(national_id="29001150101234").referred_by, referrer)

    def test_relative_link_on_registration(self):
        relative = make_patient(self.branch, nid="28501010101235", phone="01112223334")
        self.client.post("/patients/new/", self.form_data(relative_lookup=relative.file_number, relative_relation="sibling"))
        patient = Patient.objects.get(national_id="29001150101234")
        relation = PatientRelation.objects.get()
        self.assertEqual((relation.patient, relation.related_patient), (patient, relative))
        self.assertEqual(relative.relations()[0][0], patient)

    def test_convert_lead_to_patient(self):
        lead = Lead.objects.create(branch=self.branch, full_name="متصل", phone_primary="01001234567", missing_teeth="full_arch")
        page = self.client.get(f"/patients/new/?lead={lead.pk}")
        self.assertEqual(page.context["form"].initial["phone_primary"], "01001234567")
        self.client.post("/patients/new/", self.form_data(lead=lead.pk))
        lead.refresh_from_db()
        self.assertEqual(lead.status, Lead.Status.CONVERTED)
        self.assertEqual(lead.converted_patient.national_id, "29001150101234")

    def test_lead_with_phone_of_registered_patient_is_rejected(self):
        make_patient(self.branch, phone="01001234567")
        response = self.client.post("/patients/calls/new/", {
            "full_name": "x", "phone_primary": "01001234567", "preferred_phone": "primary",
            "missing_teeth": "single", "status": "new",
        })
        self.assertIn("phone_primary", response.context["form"].errors)
        self.assertFalse(Lead.objects.exists())

    def test_log_call_updates_lead(self):
        lead = Lead.objects.create(branch=self.branch, full_name="متصل", phone_primary="01001234567")
        self.client.post(f"/patients/calls/{lead.pk}/log/", {"called_at": "25/09/2026 10:00", "outcome": "no_answer"})
        lead.refresh_from_db()
        self.assertEqual(lead.status, Lead.Status.FOLLOW_UP)
        self.assertEqual(lead.calls.count(), 1)

    def test_search_by_phone_opens_the_file(self):
        patient = make_patient(self.branch)
        response = self.client.get("/patients/", {"q": "+20 100 123 4567"})
        self.assertRedirects(response, patient.get_absolute_url(), fetch_redirect_response=False)


class DentistAccessTests(TestCase):
    def setUp(self):
        self.branch = setup_clinic()
        self.dentist = make_dentist("dentist", kind="fulltime")
        self.other_dentist = make_dentist("dentist2", kind="fulltime")
        self.mine = make_patient(self.branch, assigned_dentist=self.dentist)
        self.not_mine = make_patient(self.branch, nid="28501010101235", phone="01112223334",
                                     assigned_dentist=self.other_dentist)
        self.client.login(username="dentist", password=PASSWORD)

    def test_cia_dentist_sees_every_patient_and_can_filter_their_own(self):
        # They record the work of the course candidates, who do not log in.
        self.assertEqual(self.client.get(self.not_mine.get_absolute_url()).status_code, 200)
        self.assertEqual(len(self.client.get("/patients/").context["page_obj"]), 2)
        self.assertEqual(list(self.client.get("/patients/", {"mine": "on"}).context["page_obj"]), [self.mine])

    def test_my_patients_include_patients_booked_with_them(self):
        Appointment.objects.create(branch=self.branch, patient=self.not_mine, dentist=self.dentist, scheduled_at=timezone.now())
        self.assertEqual(len(self.client.get("/patients/", {"mine": "on"}).context["page_obj"]), 2)

    def test_dentist_cannot_register_or_open_call_list(self):
        self.assertEqual(self.client.get("/patients/new/").status_code, 403)
        self.assertEqual(self.client.get("/patients/calls/").status_code, 403)

    def test_stock_manager_sees_no_patients(self):
        make_user("stock", "stock")
        self.client.login(username="stock", password=PASSWORD)
        self.assertEqual(self.client.get(self.mine.get_absolute_url()).status_code, 403)
        self.assertEqual(self.client.get("/patients/").status_code, 403)


class PlanFinderAndCallListTests(TestCase):
    def setUp(self):
        from apps.charting.models import PlanItem, TreatmentPlan
        from apps.clinical.models import TreatmentStepType

        self.branch = setup_clinic()
        self.head = make_user("head", "head_cia")
        self.secretary = make_user("sec", "secretary")
        guided = TreatmentStepType.objects.get(name_en="Guided implant surgery")
        simple = TreatmentStepType.objects.get(name_en="Implant placement")
        self.p1 = make_patient(self.branch)
        self.p2 = make_patient(self.branch, nid="28501010101235", phone="01112223334")
        self.p3 = make_patient(self.branch, nid="28501010101236", phone="01112223335")
        for patient, step_type, difficulty in [(self.p1, guided, "simple"), (self.p2, guided, "moderate"),
                                               (self.p3, simple, "simple")]:
            plan = TreatmentPlan.objects.create(patient=patient, difficulty=difficulty)
            PlanItem.objects.create(plan=plan, step_type=step_type, teeth="36")
        self.guided = guided

    def find(self, **params):
        return self.client.get("/chart/plans/", params)

    def found(self, **params):
        return {plan.patient for plan, *_rest in self.find(**params).context["page_obj"]}

    def test_find_plans_by_procedure_and_difficulty(self):
        self.client.login(username="head", password=PASSWORD)
        self.assertEqual(self.found(), {self.p1, self.p2, self.p3})
        self.assertEqual(self.found(procedures=self.guided.pk), {self.p1, self.p2})
        self.assertEqual(self.found(procedures=self.guided.pk, difficulty="simple"), {self.p1})
        Appointment.objects.create(branch=self.branch, patient=self.p1, scheduled_at=timezone.now() + timedelta(days=2))
        self.assertEqual(self.found(procedures=self.guided.pk, no_appointment="on"), {self.p2})
        csv = self.find(procedures=self.guided.pk, export="csv").content.decode("utf-8-sig")
        self.assertEqual(len(csv.strip().splitlines()), 3)
        # The reasons sent to the reception are in Arabic, whatever the sender's language.
        page = self.find(procedures=self.guided.pk)
        self.assertIn((self.p1, "زرع بدليل جراحي 36"), page.context["call_rows"])
        self.assertIn("زرع بدليل جراحي", page.context["call_title"])
        self.client.login(username="sec", password=PASSWORD)
        self.assertEqual(self.find().status_code, 403)

    def test_send_to_reception_and_record_answers(self):
        from apps.core.models import Notification
        from apps.patients.models import CallList, CallListEntry

        self.client.login(username="head", password=PASSWORD)
        response = self.client.post("/patients/to-call/new/", {
            "title": "Guided cases", "message": "Book the surgery day", "patient": [self.p1.pk, self.p2.pk],
            f"reason_{self.p1.pk}": "Guided implant surgery 36", "next": "/chart/plans/",
        })
        call_list = CallList.objects.get()
        self.assertRedirects(response, call_list.get_absolute_url(), fetch_redirect_response=False)
        self.assertEqual(call_list.entries.count(), 2)
        self.assertTrue(Notification.objects.filter(recipient=self.secretary, title__contains="Guided cases").exists())

        self.client.login(username="sec", password=PASSWORD)
        self.assertEqual(self.client.get("/").context["call_lists"][0], call_list)
        page = self.client.get(call_list.get_absolute_url())
        self.assertContains(page, "Guided implant surgery 36")
        self.assertContains(page, "الضرس الأول السفلي الأيسر")  # 36 explained to the secretary
        for entry, outcome in zip(call_list.entries.all(), ["booked", "no_answer"]):
            self.client.post(f"/patients/to-call/answer/{entry.pk}/", {f"e{entry.pk}-outcome": outcome,
                                                                       f"e{entry.pk}-response": "ok"})
        first = call_list.entries.get(patient=self.p1)
        self.assertEqual((first.outcome, first.called_by, first.attempts), ("booked", self.secretary, 1))
        self.assertEqual(call_list.progress, (2, 2))
        # The head of CIA is told when every patient was called.
        self.assertTrue(Notification.objects.filter(recipient=self.head, title__contains="Guided cases").exists())
        self.assertEqual(self.client.post("/patients/to-call/new/", {"title": "x", "patient": [self.p1.pk]}).status_code, 403)
        self.assertEqual(CallListEntry.objects.filter(outcome="pending").count(), 0)


class PatientSearchTests(TestCase):
    def setUp(self):
        self.branch = setup_clinic()
        self.patient = make_patient(self.branch, name="أحمد فاطمة مصطفى", nid="29001011234567", phone="01001234567")
        make_patient(self.branch, name="سارة علي", nid="29002021234568", phone="01101234567")
        make_user("sec", "secretary")
        self.client.login(username="sec", password=PASSWORD)

    def test_suggestions_ignore_arabic_spelling_variants(self):
        for typed in ("احمد", "فاطمه مصطفي", "0100123", self.patient.file_number):
            results = self.client.get("/patients/lookup/", {"q": typed}).json()["results"]
            self.assertEqual([r["value"] for r in results], [self.patient.file_number], typed)
        self.assertEqual(self.client.get("/patients/lookup/", {"q": "x"}).json()["results"], [])

    def test_booking_takes_the_chosen_patient_and_a_dd_mm_yyyy_date(self):
        from apps.scheduling.models import Appointment

        day = timezone.localdate() + timedelta(days=2)
        response = self.client.post("/schedule/appointments/new/", {
            "patient_lookup": self.patient.file_number, "scheduled_at_0": day.strftime("%d/%m/%Y"),
            "scheduled_at_1": "10:15", "duration_minutes": "30",
        })
        appointment = Appointment.objects.get()
        self.assertRedirects(response, appointment.get_absolute_url(), fetch_redirect_response=False)
        local = timezone.localtime(appointment.scheduled_at)
        self.assertEqual((local.date(), local.strftime("%H:%M"), appointment.patient), (day, "10:15", self.patient))
        page = self.client.get(f"/schedule/appointments/{appointment.pk}/edit/").content.decode()
        self.assertIn(day.strftime("%d/%m/%Y"), page)
        self.assertIn('value="10:15" selected', page)


class TreatmentExplanationTests(TestCase):
    def test_the_patient_file_explains_plans_and_treatments_in_arabic(self):
        from apps.charting.models import PlanItem, TreatmentPlan
        from apps.charting.teeth import teeth_ar
        from apps.clinical.models import TreatmentStep, TreatmentStepType

        self.assertEqual(teeth_ar("11-12"), "12 (القاطع الجانبي العلوي الأيمن)، 11 (القاطع الأوسط العلوي الأيمن)")
        branch = setup_clinic()
        patient = make_patient(branch)
        guided = TreatmentStepType.objects.get(name_en="Guided implant surgery")
        plan = TreatmentPlan.objects.create(patient=patient)
        PlanItem.objects.create(plan=plan, step_type=guided, teeth="46")
        scaling = TreatmentStepType.objects.get(name_en="Scaling")
        TreatmentStep.objects.create(patient=patient, step_type=scaling, performed_at=timezone.now())
        make_user("sec", "secretary")
        self.client.login(username="sec", password=PASSWORD)
        # The dental work is hidden from the reception until the owner shows it (Settings → Access).
        self.assertNotContains(self.client.get(f"/patients/{patient.pk}/"), "زرع بدليل جراحي")
        from apps.core.models import ClinicSettings

        options = ClinicSettings.get()
        options.reception_sees = ["plan", "steps"]
        options.save()
        page = self.client.get(f"/patients/{patient.pk}/")
        for text in ("زرع بدليل جراحي", guided.description_ar, "الضرس الأول السفلي الأيمن", scaling.description_ar):
            self.assertContains(page, text)
        self.assertNotContains(page, "/clinical/steps/")  # the secretary reads it here, without the treatment log


class PatientDataTests(TestCase):
    def setUp(self):
        self.branch = setup_clinic()
        make_user("sec", "secretary")
        self.client.login(username="sec", password=PASSWORD)
        self.source = ReferralSource.objects.get(name_en="Facebook")

    def test_old_files_keep_their_date_and_the_governorate_comes_from_the_id(self):
        self.client.post("/patients/new/", {
            "full_name": "محمد أحمد علي", "id_type": "nid", "national_id": "29001152101234", "phone_primary": "01001234567",
            "preferred_phone": "primary", "missing_teeth": "single", "referral_source": self.source.pk,
            "registered_on": "03/02/2019", "phone_secondary": "01101234567", "marital_status": "married",
            "occupation": "محاسب", "city": "الدقي",
        })
        patient = Patient.objects.get()
        self.assertEqual((patient.registered_on, patient.governorate), (date(2019, 2, 3), "21"))  # Giza
        self.assertEqual(Patient._meta.get_field("governorate").choices[0][0], "01")

    def test_out_needs_a_reason(self):
        from apps.patients.models import OutReason

        patient = make_patient(self.branch)
        data = {"full_name": patient.full_name, "id_type": "nid", "national_id": patient.national_id,
                "phone_primary": patient.phone_primary, "preferred_phone": "primary", "missing_teeth": "single",
                "referral_source": self.source.pk, "status": "out"}
        response = self.client.post(f"/patients/{patient.pk}/edit/", data)
        self.assertIn("out_reason", response.context["form"].errors)
        self.assertTrue(OutReason.objects.exists())

    def test_the_call_list_keeps_who_called_first(self):
        for name, phone, day in [("Late caller", "01001111111", ""), ("Early caller", "01002222222", "01/09/2026")]:
            self.client.post("/patients/calls/new/", {"full_name": name, "phone_primary": phone,
                                                      "preferred_phone": "primary", "missing_teeth": "single",
                                                      "status": "new", "first_call_on": day})
        late = Lead.objects.get(full_name="Late caller")
        self.assertEqual(late.first_call_on, timezone.localdate())
        page = self.client.get("/patients/calls/", {"due": "1"})
        self.assertEqual([lead.full_name for lead in page.context["page_obj"]], ["Early caller", "Late caller"])


class PhoneCheckTests(TestCase):
    def test_wrong_length_and_duplicates_are_told_while_typing(self):
        branch = setup_clinic()
        patient = make_patient(branch, name="سارة علي", phone="01001234567")
        make_user("sec", "secretary")
        self.client.login(username="sec", password=PASSWORD)

        def check(**params):
            return self.client.get("/patients/phone-check/", params).json()

        self.assertIn("10", check(phone="0100123456")["message"])
        self.assertIn("سارة علي", check(phone="01001234567")["message"])
        self.assertEqual(check(phone="01001234567", pk=patient.pk), {})  # its own number
        self.assertEqual(check(phone="01101234567"), {})
        self.assertEqual(check(phone="0223456789", field="secondary"), {})  # a landline is fine as mobile 2
        response = self.client.post("/patients/new/", {"full_name": "x", "phone_primary": "0100"})
        self.assertContains(response, "data-show-on-load")  # the errors open in a box


@override_settings(MEDIA_ROOT=MEDIA)
class IdCardTests(TestCase):
    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        shutil.rmtree(MEDIA, ignore_errors=True)

    @staticmethod
    def photo(size=(1200, 1000), card=(300, 250, 900, 628), name="id.jpg"):
        """A grey table with a light blue card on it, as a phone photo would show."""
        import io

        from PIL import Image, ImageDraw

        image = Image.new("RGB", size, (90, 90, 95))
        draw = ImageDraw.Draw(image)
        draw.rectangle(card, fill=(200, 225, 240))
        draw.text((card[0] + 40, card[1] + 40), "29001011234567", fill=(0, 0, 0))
        output = io.BytesIO()
        image.save(output, "JPEG")
        return SimpleUploadedFile(name, output.getvalue(), content_type="image/jpeg")

    def test_the_card_is_cut_out_and_the_original_kept(self):
        from PIL import Image

        branch = setup_clinic()
        patient = make_patient(branch)
        document = PatientDocument.objects.create(patient=patient, kind=PatientDocument.Kind.ID_FRONT, file=self.photo())
        with Image.open(document.file.path) as card:
            self.assertTrue(590 <= card.width <= 640 and 360 <= card.height <= 410, card.size)
        self.assertTrue(document.original)
        # A photo that is already just the card is kept as it is.
        plain = PatientDocument.objects.create(patient=patient, kind=PatientDocument.Kind.ID_BACK,
                                               file=self.photo(size=(856, 540), card=(0, 0, 856, 540)))
        self.assertFalse(plain.original)
        # A card photographed standing up is turned to lie flat.
        standing = PatientDocument.objects.create(patient=patient, kind=PatientDocument.Kind.ID_BACK,
                                                  file=self.photo(size=(1000, 1200), card=(250, 300, 628, 900)))
        with Image.open(standing.file.path) as card:
            self.assertGreater(card.width, card.height)
        make_user("sec", "secretary")
        self.client.login(username="sec", password=PASSWORD)
        self.client.post(f"/patients/{patient.pk}/documents/{document.pk}/rotate/", {"way": "right"})
        document.refresh_from_db()
        with Image.open(document.file.path) as card:
            self.assertGreater(card.height, card.width)


@override_settings(MEDIA_ROOT=MEDIA)
class XrayAndCbctTests(TestCase):
    """X-rays may be bigger than other documents, and a CBCT (too big to upload) is kept as where it is."""

    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        shutil.rmtree(MEDIA, ignore_errors=True)

    def setUp(self):
        branch = setup_clinic()
        self.dentist = make_dentist("dentist", kind="candidate")
        self.patient = make_patient(branch, assigned_dentist=self.dentist)
        make_user("sec", "secretary")
        self.client.login(username="sec", password=PASSWORD)
        self.url = f"/patients/{self.patient.pk}/documents/"

    def test_a_cbct_folder_or_link_is_kept_and_shown_to_the_dentist(self):
        self.client.post(self.url, {"kind": "xray", "location": r"\\CIA-SERVER\CBCT\CIA-00001", "notes": "CBCT lower"})
        self.client.post(self.url, {"kind": "xray", "location": "https://viewer.example.com/study/42"})
        self.client.post(self.url, {"kind": "xray", "location": "javascript:alert(1)"})
        self.assertEqual(PatientDocument.objects.filter(kind="xray").count(), 3)
        self.client.post(self.url, {"kind": "other"})  # nothing to keep
        self.assertEqual(PatientDocument.objects.count(), 3)
        self.client.login(username="dentist", password=PASSWORD)
        page = self.client.get(f"/patients/{self.patient.pk}/")
        self.assertContains(page, r"\\CIA-SERVER\CBCT\CIA-00001")
        self.assertContains(page, 'data-copy="\\\\CIA-SERVER\\CBCT\\CIA-00001"')
        self.assertContains(page, 'href="https://viewer.example.com/study/42"')
        self.assertNotContains(page, 'href="javascript:')  # only web links become links
        self.assertEqual(len(page.context["xrays"]), 3)

    @override_settings(MAX_UPLOAD_SIZE_MB=1, MAX_XRAY_UPLOAD_MB=3)
    def test_xrays_may_be_bigger_than_other_documents(self):
        big = b"%PDF-1.4" + b"0" * (2 * 1024 * 1024)
        self.client.post(self.url, {"kind": "other", "file": SimpleUploadedFile("scan.pdf", big, "application/pdf")})
        self.assertFalse(PatientDocument.objects.exists())
        self.client.post(self.url, {"kind": "xray", "file": SimpleUploadedFile("opg.pdf", big, "application/pdf")})
        self.assertTrue(PatientDocument.objects.filter(kind="xray").exists())


class PatientsPerPlaceTests(TestCase):
    """Each place has its own patients; moving to another place opens a new file there."""

    def setUp(self):
        from apps.core.models import Branch, UserProfile

        self.cia = setup_clinic()
        self.cic = Branch.objects.get(code="CIC")
        self.secretary = make_user("sec", "secretary")
        UserProfile.objects.update_or_create(user=self.secretary, defaults={"branch": self.cia})
        self.secretary.profile.places.set([self.cia, self.cic])
        self.patient = make_patient(self.cia, name="مريض الأكاديمية الأول")
        self.client.login(username="sec", password=PASSWORD)

    def work_at(self, code):
        self.client.post("/place/", {"place": code, "next": "/"})

    def test_a_place_does_not_see_the_patients_of_another_place(self):
        self.assertEqual(len(self.client.get("/patients/lookup/?q=مريض").json()["results"]), 1)
        self.work_at("CIC")
        self.assertEqual(self.client.get("/patients/lookup/?q=مريض").json()["results"], [])
        self.assertNotContains(self.client.get("/patients/?q=مريض"), "مريض الأكاديمية الأول")
        other = self.client.get(self.patient.get_absolute_url())  # she works at CIA too: switch first
        self.assertContains(other, "العمل في CIA وفتح الملف")
        self.assertNotContains(other, "مريض الأكاديمية الأول")
        # The same person cannot get a second file at CIC; the message gives the CIA file number, not the name.
        response = self.client.post("/patients/new/", {"full_name": "مريض الأكاديمية الأول", "id_type": "nid",
                                                       "national_id": self.patient.national_id,
                                                       "phone_primary": "01005556666"})
        self.assertIn(self.patient.file_number, str(response.context["form"].errors["national_id"]))
        self.assertNotIn("مريض الأكاديمية", str(response.context["form"].errors["national_id"]))

    def test_moving_to_another_place_opens_a_new_file_and_closes_this_one(self):
        from datetime import datetime, time

        from apps.scheduling.models import Appointment

        coming = Appointment.objects.create(branch=self.cia, patient=self.patient, scheduled_at=timezone.make_aware(
            datetime.combine(timezone.localdate() + timedelta(days=3), time(12))))
        page = self.client.get(self.patient.get_absolute_url())
        self.assertContains(page, "نقل لمكان آخر")
        self.assertNotContains(page, "Book at")  # booking at another place is not done
        self.client.post(f"/patients/{self.patient.pk}/move/", {"place": "CIC"})
        self.patient.refresh_from_db()
        moved = self.patient.transferred_to
        self.assertEqual(self.patient.status, Patient.Status.OUT)
        self.assertEqual((moved.branch, moved.full_name, moved.national_id), (self.cic, self.patient.full_name,
                                                                               self.patient.national_id))
        self.assertTrue(moved.file_number.startswith("CIC-"))
        coming.refresh_from_db()
        self.assertEqual(coming.status, Appointment.Status.CANCELLED)
        self.assertContains(self.client.get(self.patient.get_absolute_url()), moved.file_number)
        # Coming back to CIA opens the old CIA file again, not a third one.
        self.work_at("CIC")
        self.client.post(f"/patients/{moved.pk}/move/", {"place": "CIA"})
        self.patient.refresh_from_db()
        moved.refresh_from_db()
        self.assertEqual((self.patient.status, moved.status), (Patient.Status.ACTIVE, Patient.Status.OUT))
        self.assertEqual(Patient.objects.count(), 2)

    def test_complaints_follow_the_place(self):
        from apps.complaints.models import Complaint

        Complaint.objects.create(branch=self.cia, patient=self.patient, category="other", description="الانتظار")
        self.assertEqual(len(self.client.get("/complaints/?status=").context["object_list"]), 1)
        self.work_at("CIC")
        self.assertEqual(len(self.client.get("/complaints/?status=").context["object_list"]), 0)


class ArabicNameTests(TestCase):
    def test_the_name_is_in_arabic_with_at_least_three_names(self):
        from django import forms as django_forms

        from apps.patients.forms import clean_arabic_name

        self.assertEqual(clean_arabic_name("  محمد   أحمد علي "), "محمد أحمد علي")
        for wrong in ("Mohamed Ahmed Ali", "محمد أحمد", "محمد أحمد 12", "محمد Ahmed علي"):
            with self.assertRaises(django_forms.ValidationError):
                clean_arabic_name(wrong)
        setup_clinic()
        make_user("sec", "secretary")
        self.client.login(username="sec", password=PASSWORD)
        response = self.client.post("/patients/new/", {"full_name": "Ahmed Ali Hassan"})
        self.assertIn("العربية", str(response.context["form"].errors["full_name"]))


class FileStepsTests(TestCase):
    """The file is filled in order: medical history, dental history, examination, plan, surgery chart."""

    def setUp(self):
        self.branch = setup_clinic()
        self.patient = make_patient(self.branch)
        make_dentist("dentist", kind="fulltime")
        make_user("sec", "secretary")
        self.client.login(username="dentist", password=PASSWORD)

    def steps(self):
        return {s["code"]: s for s in self.client.get(f"/patients/{self.patient.pk}/").context["file_steps"]}

    def test_the_dentist_goes_through_the_file_in_order(self):
        from apps.charting.models import Examination

        steps = self.steps()
        self.assertTrue(steps["medical"]["next"])
        self.assertTrue(steps["surgery"]["optional"])
        page = self.client.get(steps["medical"]["url"])
        self.assertNotIn("smoker", page.context["form"].fields)  # the dental history is the next page
        self.assertIn("conditions", page.context["form"].fields)
        response = self.client.post(steps["medical"]["url"], {"bp_last_systolic": "130", "allergy_penicillin": "on"})
        self.assertRedirects(response, f"/patients/{self.patient.pk}/history/?part=dental&flow=1")
        history = Examination.objects.get()
        self.assertEqual((history.medical_taken, history.dental_taken, history.history_only), (True, False, True))
        form = self.client.get(response.url).context["form"]
        self.assertIn("cooperation_score", form.fields)  # the dentist's own judgement, asked of the dentist only
        response = self.client.post(response.url, {"smoker": "on", "cigarettes_per_day": "10"})
        self.assertRedirects(response, f"/chart/patient/{self.patient.pk}/exam/new/?flow=1", fetch_redirect_response=False)
        history.refresh_from_db()
        self.assertEqual((history.dental_taken, history.bp_last_systolic, history.smoker), (True, 130, True))
        self.assertEqual(Examination.objects.count(), 1)
        steps = self.steps()
        self.assertTrue(steps["exam"]["next"] and steps["medical"]["done"] and steps["dental"]["done"])

    def test_the_reception_sees_a_summary_and_what_the_owner_allows(self):
        from apps.charting.models import Examination
        from apps.core.models import ClinicSettings
        from apps.patients.models import MedicalCondition

        exam = Examination.objects.create(patient=self.patient, allergy_penicillin=True, bp_last_systolic=150,
                                          bp_last_diastolic=95)
        exam.conditions.add(MedicalCondition.objects.get(name_en="Diabetes"))
        self.client.login(username="sec", password=PASSWORD)
        page = self.client.get(f"/patients/{self.patient.pk}/")
        self.assertIsNone(page.context.get("file_steps"))
        self.assertContains(page, "medical-summary")
        self.assertNotContains(page, "150/95")  # the full history stays with the dentist
        self.assertNotContains(page, 'data-bs-target="#plan"')
        self.assertContains(page, 'data-bs-target="#lab"')  # shown until the owner hides it
        self.assertEqual(self.client.get(f"/patients/{self.patient.pk}/history/").status_code, 403)
        owner = make_user("boss", "owner")
        self.client.force_login(owner)
        self.client.post("/settings/access/", {"what": "file_parts", "parts": ["medical", "plan"]})
        self.assertEqual(ClinicSettings.get().reception_sees, ["medical", "plan"])
        self.client.login(username="sec", password=PASSWORD)
        page = self.client.get(f"/patients/{self.patient.pk}/")
        self.assertContains(page, "150/95")
        self.assertContains(page, 'data-bs-target="#plan"')
        self.assertNotContains(page, 'data-bs-target="#lab"')

    def test_the_file_as_excel(self):
        from openpyxl import load_workbook

        response = self.client.get(f"/patients/{self.patient.pk}/excel/")
        self.assertEqual(response.status_code, 200)
        book = load_workbook(io.BytesIO(b"".join(response.streaming_content)))
        self.assertIn("Personal data", book.sheetnames)
        self.assertIn(self.patient.file_number, [row[1] for row in book["Personal data"].iter_rows(values_only=True)])
        self.client.login(username="sec", password=PASSWORD)
        self.assertEqual(self.client.get(f"/patients/{self.patient.pk}/excel/").status_code, 403)


class MedicalFollowUpTests(TestCase):
    """Round 10: readings above the limits put the patient on the medical follow-up; the dentist writes a ready
    consultation letter; the answer is recorded and the dentist is told."""

    def setUp(self):
        from decimal import Decimal

        from apps.charting.models import Examination

        self.branch = setup_clinic()
        self.patient = make_patient(self.branch)
        self.dentist = make_dentist("dentist", kind="fulltime")
        self.patient.assigned_dentist = self.dentist
        self.patient.save()
        make_user("sec", "secretary")
        self.exam = Examination.objects.create(patient=self.patient, history_only=True, hba1c=Decimal("8.2"),
                                               bp_clinic_systolic=150, bp_clinic_diastolic=90, allergy_penicillin=True)
        self.client.login(username="dentist", password=PASSWORD)

    def test_the_readings_put_the_patient_on_the_list(self):
        from apps.patients.medical import clearance, reading_flags

        self.assertEqual([code for code, _text in reading_flags(self.exam)], ["hba1c"])  # 150/90 is under 160/100
        self.assertEqual(clearance(self.patient)["state"], "needed")
        page = self.client.get("/patients/medical-follow-up/?mine=0")
        self.assertEqual([exam.patient for exam, _flags in page.context["needing"]], [self.patient])
        self.assertContains(self.client.get(f"/patients/{self.patient.pk}/"), "clearance-box")
        # Another place never sees this patient.
        from apps.core.models import Branch

        other = make_patient(Branch.objects.get(code="CIC"), name="مريض آخر", nid="29001011234568", phone="01001234568")
        self.assertNotIn(other, [e.patient for e, _f in page.context["needing"]])

    def test_the_letter_is_ready_made(self):
        page = self.client.get(f"/patients/{self.patient.pk}/consult/")
        form = page.context["form"]
        self.assertIn("Artinibsa 4%", form.initial.get("anesthesia") or form.instance.anesthesia)
        self.assertEqual(form.initial["reasons"], ["hba1c"])
        self.assertIn("HbA1c 8.2%", form.initial["findings"])
        self.assertIn("Clindamycin", form.initial["medications"])  # allergic to penicillin
        medicines = self.client.get(f"/patients/{self.patient.pk}/consult/medicines/?procedure=sinus").json()
        self.assertIn("Xylometazoline", medicines["medications"])
        response = self.client.post(f"/patients/{self.patient.pk}/consult/", {
            "reasons": ["hba1c"], "findings": form.initial["findings"], "physician": "Adel Sami",
            "specialty": "Endocrinology (diabetes)", "procedures": ["implants", "sinus"], "procedure_details": "16",
            "duration": "About 1 hour", "bleeding": "minor", "anesthesia": form.instance.anesthesia,
            "medications": "Clindamycin 300 mg", "question": "Fit?", "sent_on": "01/09/2026",
            "dentist": self.dentist.pk})
        from apps.patients.models import MedicalConsult

        consult = MedicalConsult.objects.get()
        self.assertRedirects(response, consult.get_absolute_url(), fetch_redirect_response=False)
        self.assertEqual(consult.bleeding, "moderate")  # a sinus lift
        self.assertEqual(consult.procedure_labels(), ["Implant placement", "Sinus lift"])
        page = self.client.get(consult.get_absolute_url())
        self.assertContains(page, "Dr. Adel Sami")
        self.assertContains(page, "Medical consultation before dental surgery")
        from apps.patients.medical import clearance

        self.assertEqual(clearance(self.patient)["state"], "waiting")
        # The reception copies the answer the patient brought back; the dentist is told.
        self.client.login(username="sec", password=PASSWORD)
        self.assertEqual(self.client.get(f"/patients/{self.patient.pk}/consult/").status_code, 403)
        response = self.client.post(f"/patients/consults/{consult.pk}/answer/", {"answer": "precautions"})
        self.assertEqual(response.status_code, 200)  # the precautions must be written
        self.client.post(f"/patients/consults/{consult.pk}/answer/", {
            "answer": "precautions", "answer_notes": "Stop aspirin 5 days before", "answered_by": "Dr. Adel Sami",
            "answered_on": "05/09/2026"})
        consult.refresh_from_db()
        self.assertEqual((consult.status, consult.answer), ("answered", "precautions"))
        self.assertEqual(clearance(self.patient)["state"], "fit")
        from apps.core.models import Notification

        self.assertTrue(Notification.objects.filter(recipient=self.dentist.user, url=consult.get_absolute_url()).exists())

    def test_not_needed_and_postponed(self):
        from apps.patients.medical import clearance
        from apps.patients.models import MedicalConsult

        self.client.post(f"/patients/{self.patient.pk}/consult/not-needed/", {"note": "HbA1c of last year"})
        self.assertEqual(clearance(self.patient)["state"], "cleared")
        page = self.client.get("/patients/medical-follow-up/?mine=0")
        self.assertEqual(list(page.context["needing"]), [])
        MedicalConsult.objects.all().delete()
        held = MedicalConsult.objects.create(patient=self.patient, dentist=self.dentist, status="answered",
                                             answer="postpone", recheck_on=timezone.localdate())
        self.assertEqual(clearance(self.patient)["state"], "not_fit")
        page = self.client.get("/patients/medical-follow-up/?mine=0")
        self.assertEqual(page.context["held"], [held])


class JourneyTests(TestCase):
    """Round 10: the file follows the dentist's day: histories and readings, examination, impression or scan,
    CBCT, planning and plan, then surgery, restorative work and delivery."""

    def setUp(self):
        self.branch = setup_clinic()
        self.patient = make_patient(self.branch)
        make_dentist("dentist", kind="fulltime")
        make_user("sec", "secretary")
        self.client.login(username="dentist", password=PASSWORD)

    def steps(self):
        return {s["code"]: s for s in self.client.get(f"/patients/{self.patient.pk}/").context["file_steps"]}

    def test_the_steps_in_the_dentists_order(self):
        from apps.charting.models import Examination

        codes = list(self.steps())
        self.assertEqual(codes, ["medical", "dental", "exam", "impression", "cbct", "plan", "surgery"])
        page = self.client.get(f"/patients/{self.patient.pk}/history/?part=medical&flow=1")
        fields = list(page.context["form"].fields)
        self.assertEqual(fields[:3], ["bp_clinic_systolic", "bp_clinic_diastolic", "glucose_random_clinic"])
        Examination.objects.create(patient=self.patient, history_only=True)
        Examination.objects.create(patient=self.patient, teeth_missing="36")
        steps = self.steps()
        self.assertTrue(steps["impression"]["next"])
        self.assertEqual(self.client.get(f"/patients/{self.patient.pk}/").context["next_step"]["code"], "impression")
        # The examination page folds the histories taken before.
        self.assertTrue(self.client.get(f"/chart/patient/{self.patient.pk}/exam/new/").context["form"].folded)

    def test_impression_and_cbct(self):
        from apps.charting.models import Examination
        from apps.clinical.models import TreatmentStep, TreatmentStepType

        Examination.objects.create(patient=self.patient, teeth_missing="36")
        scan = TreatmentStepType.objects.get(name_en="Diagnostic intraoral scan")
        page = self.client.get(f"/patients/{self.patient.pk}/records/?step=impression&flow=1")
        self.assertIn(scan, list(page.context["types"]))
        response = self.client.post(f"/patients/{self.patient.pk}/records/?step=impression&flow=1",
                                    {"type": scan.pk, "flow": "1"})
        self.assertRedirects(response, f"/patients/{self.patient.pk}/records/?step=cbct&flow=1",
                             fetch_redirect_response=False)
        self.assertEqual(TreatmentStep.objects.get().step_type, scan)
        self.assertTrue(self.steps()["impression"]["done"])
        response = self.client.post(f"/patients/{self.patient.pk}/records/?step=cbct&flow=1", {"type": "skip",
                                                                                              "flow": "1"})
        self.assertRedirects(response, f"/chart/patient/{self.patient.pk}/plan/new/?flow=1",
                             fetch_redirect_response=False)
        self.client.login(username="sec", password=PASSWORD)
        self.assertEqual(self.client.get(f"/patients/{self.patient.pk}/records/").status_code, 403)

    def test_later_steps_show_when_they_are_needed(self):
        from decimal import Decimal

        from apps.charting.models import Examination, PlanItem, TreatmentPlan
        from apps.clinical.models import TreatmentStepType

        Examination.objects.create(patient=self.patient, hba1c=Decimal("9.1"))
        plan = TreatmentPlan.objects.create(patient=self.patient)
        PlanItem.objects.create(plan=plan, step_type=TreatmentStepType.objects.get(name_en="Composite restoration"),
                                teeth="16")
        steps = self.steps()
        self.assertIn("fitness", steps)
        self.assertIn("restorative", steps)
        self.assertTrue(steps["impression"]["skipped"])  # the plan is written: the impression was not needed
        self.assertEqual(steps["restorative"]["note"], "1 left")

    def test_registration_folds_the_history(self):
        self.client.login(username="sec", password=PASSWORD)
        page = self.client.get("/patients/new/")
        self.assertContains(page, 'class="form-fold')


@override_settings(MEDIA_ROOT=MEDIA)
class Round13RegistrationTests(TestCase):
    """The ID photos kept after an error, the details required at registration, the city list, the visit
    preferences and the referring patient made optional."""

    def setUp(self):
        self.branch = setup_clinic()
        self.secretary = make_user("sec", "secretary")
        self.client.login(username="sec", password=PASSWORD)
        self.source = ReferralSource.objects.get(name_en="Facebook")

    def data(self, **overrides):
        data = {"full_name": "محمد أحمد علي حسن", "id_type": "nid", "national_id": "29001152101234",
                "phone_primary": "01001234567", "phone_secondary": "01101234567", "preferred_phone": "primary",
                "missing_teeth": "single", "referral_source": self.source.pk, "marital_status": "married",
                "occupation": "مهندس", "city": "الدقي"}
        data.update(overrides)
        return data

    @staticmethod
    def photo(name="id.jpg"):
        return SimpleUploadedFile(name, b"\xff\xd8\xff fake-image", content_type="image/jpeg")

    def test_mobile_2_marital_status_occupation_and_city_are_required(self):
        response = self.client.post("/patients/new/", self.data(phone_secondary="", marital_status="", occupation="",
                                                                city=""))
        errors = response.context["form"].errors
        for name in ("phone_secondary", "marital_status", "occupation", "city"):
            self.assertIn(name, errors, name)
        self.assertFalse(Patient.objects.exists())

    def test_birth_date_gender_and_governorate_come_from_the_id_or_are_required(self):
        self.client.post("/patients/new/", self.data())
        patient = Patient.objects.get()
        self.assertEqual((patient.birth_date, patient.gender, patient.governorate), (date(1990, 1, 15), "M", "21"))
        response = self.client.post("/patients/new/", self.data(id_type="passport", national_id="A1234567",
                                                                phone_primary="01001234568",
                                                                phone_secondary="01101234568"))
        errors = response.context["form"].errors
        for name in ("birth_date", "gender", "governorate"):
            self.assertIn(name, errors, name)

    def test_the_id_photos_are_kept_when_the_form_comes_back_with_an_error(self):
        response = self.client.post("/patients/new/", self.data(occupation="", id_front=self.photo(),
                                                                id_back=self.photo("back.jpg")))
        self.assertIn("occupation", response.context["form"].errors)
        tokens = {name: response.context["form"].fields[name].widget.kept_token for name in ("id_front", "id_back")}
        self.assertTrue(all(tokens.values()))
        self.assertContains(response, "id_front_kept")
        # The kept photo opens for this person only.
        self.assertEqual(self.client.get(f"/kept/{tokens['id_front']}/").status_code, 200)
        other = self.client_class()
        make_user("sec2", "secretary")
        other.login(username="sec2", password=PASSWORD)
        self.assertEqual(other.get(f"/kept/{tokens['id_front']}/").status_code, 404)
        # The next save takes them without choosing them again.
        self.client.post("/patients/new/", self.data(id_front_kept=tokens["id_front"], id_back_kept=tokens["id_back"]))
        patient = Patient.objects.get()
        self.assertEqual(sorted(patient.documents.values_list("kind", flat=True)), ["id_back", "id_front"])
        self.assertEqual(self.client.get(f"/kept/{tokens['id_front']}/").status_code, 404)  # used once

    def test_id_photos_can_be_added_from_the_edit_page(self):
        patient = make_patient(self.branch, nid="29001152101234")
        data = self.data(status="active", id_front=self.photo())
        self.client.post(f"/patients/{patient.pk}/edit/", data)
        self.assertEqual(list(patient.documents.values_list("kind", flat=True)), ["id_front"])

    def test_visit_preferences_and_how_far_he_lives(self):
        self.client.post("/patients/new/", self.data(travel_minutes="45", preferred_days=["5", "1"],
                                                     preferred_times=["evening", "morning"]))
        patient = Patient.objects.get()
        self.assertEqual((patient.travel_minutes, patient.preferred_days, patient.preferred_times),
                         (45, "5,1", "morning,evening"))
        saturday_morning = timezone.make_aware(timezone.datetime(2026, 10, 3, 10, 0))
        sunday_morning = timezone.make_aware(timezone.datetime(2026, 10, 4, 10, 0))
        saturday_noon = timezone.make_aware(timezone.datetime(2026, 10, 3, 13, 0))
        self.assertTrue(patient.prefers(saturday_morning))
        self.assertFalse(patient.prefers(sunday_morning))
        self.assertFalse(patient.prefers(saturday_noon))
        answer = self.client.get(f"/patients/prefers/?q={patient.file_number}").json()
        self.assertEqual(answer["days"], [5, 1])
        self.assertEqual(answer["times"], [[9, 12], [18, 21]])
        self.assertIn("45", answer["text"])
        self.assertContains(self.client.get(patient.get_absolute_url()), "45")

    def test_the_city_list_follows_the_governorate(self):
        from apps.core.egypt import CITIES, city_in

        page = self.client.get("/patients/new/").content.decode()
        self.assertIn("data-city-list", page)
        self.assertIn("مدينة نصر", CITIES["01"])
        self.assertEqual(city_in("12 ش التحرير - الدقى - الجيزة", "21"), "الدقي")
        self.assertEqual(city_in("مدينه نصر اول", "01"), "مدينة نصر")
        self.assertEqual(city_in("لا شيء هنا", "01"), "")


class Round13PlanFinderPartsTests(TestCase):
    """The plan finder has two pages: the implant and surgery part of the plans, and the restorative part."""

    setUp = PlanFinderAndCallListTests.setUp

    def test_implant_and_restorative_plans_on_their_own_pages(self):
        from apps.charting.models import PlanItem, TreatmentPlan
        from apps.clinical.models import TreatmentStepType

        crown = TreatmentStepType.objects.filter(category=TreatmentStepType.Category.RESTORATIVE).first()
        p4 = make_patient(self.branch, nid="28501010101237", phone="01112223336")
        plan = TreatmentPlan.objects.create(patient=p4)
        PlanItem.objects.create(plan=plan, step_type=crown, teeth="11")
        self.client.login(username="head", password=PASSWORD)
        implant = self.client.get("/chart/plans/")
        self.assertEqual({p.patient for p, *_rest in implant.context["page_obj"]}, {self.p1, self.p2, self.p3})
        self.assertNotIn(crown, implant.context["form"].fields["procedures"].queryset)
        restorative = self.client.get("/chart/plans/restorative/")
        self.assertEqual({p.patient for p, *_rest in restorative.context["page_obj"]}, {p4})
        self.assertIn(crown, restorative.context["form"].fields["procedures"].queryset)
        self.assertNotIn(self.guided, restorative.context["form"].fields["procedures"].queryset)
        self.assertContains(restorative, 'href="/chart/plans/"')  # the tabs between the finders


@override_settings(MEDIA_ROOT=MEDIA)
class Round15PatientTests(TestCase):
    """Round 15: the old paper files typed in with their number, the travel time in hours and minutes, and the
    expected patients booked from the call list."""

    def setUp(self):
        self.branch = setup_clinic()
        self.secretary = make_user("sec", "secretary")
        self.client.login(username="sec", password=PASSWORD)
        self.facebook = ReferralSource.objects.get(name_en="Facebook")

    def test_old_paper_file_keeps_its_number_and_needs_less(self):
        response = self.client.post("/patients/new/", {
            "old_file": "on", "typed_file_number": "1234", "full_name": "محمد أحمد علي حسن", "id_type": "nid",
            "national_id": "29001150101234", "phone_primary": "01001234567", "preferred_phone": "primary",
            "missing_teeth": "single", "travel_minutes_0": "1", "travel_minutes_1": "30",
        })
        patient = Patient.objects.get()
        self.assertRedirects(response, f"/patients/{patient.pk}/", fetch_redirect_response=False)
        self.assertEqual(patient.file_number, "CIA-01234")
        self.assertEqual(patient.travel_minutes, 90)
        # The next file skips a number already typed by hand.
        Patient.objects.filter(pk=patient.pk).update(file_number=f"CIA-{patient.pk + 1:05d}")
        other = make_patient(self.branch, name="مريض تجربة ثاني", nid="29001150101235", phone="01001234568")
        self.assertEqual(other.file_number, f"CIA-{other.pk + 1:05d}")

    def test_typed_number_must_be_free_and_old_file_needs_one(self):
        make_patient(self.branch, file_number="CIA-00077")
        response = self.client.post("/patients/new/", {
            "old_file": "on", "typed_file_number": "77", "full_name": "محمد أحمد علي حسن", "id_type": "nid",
            "national_id": "29001150101234", "phone_primary": "01001234567", "preferred_phone": "primary",
            "missing_teeth": "single"})
        self.assertContains(response, "رقم الملف CIA-00077 مستخدم")  # the secretary reads Arabic
        response = self.client.post("/patients/new/", {
            "old_file": "on", "full_name": "محمد أحمد علي حسن", "id_type": "nid", "national_id": "29001150101234",
            "phone_primary": "01001234567", "preferred_phone": "primary", "missing_teeth": "single"})
        self.assertContains(response, "اكتبي رقم الملف الورقي.")
        # Without the tick, the round 13 boxes are still needed.
        response = self.client.post("/patients/new/", {
            "full_name": "محمد أحمد علي حسن", "id_type": "nid", "national_id": "29001150101234",
            "phone_primary": "01001234567", "preferred_phone": "primary", "missing_teeth": "single"})
        self.assertEqual(Patient.objects.count(), 1)
        self.assertIn("phone_secondary", response.context["form"].errors)

    def test_travel_time_shows_hours_and_minutes(self):
        patient = make_patient(self.branch, travel_minutes=75)
        response = self.client.get(f"/patients/{patient.pk}/")
        self.assertContains(response, "1 س 15 د")  # the secretary reads Arabic

    def test_expected_patient_booked_then_completed(self):
        lead = Lead.objects.create(branch=self.branch, full_name="سعيد محمود علي", phone_primary="01005556667",
                                   referral_source=self.facebook)
        response = self.client.post(f"/patients/calls/{lead.pk}/book/")
        patient = Patient.objects.get()
        self.assertRedirects(response, f"/schedule/appointments/new/?patient={patient.pk}",
                             fetch_redirect_response=False)
        self.assertTrue(patient.is_expected)
        self.assertEqual(patient.national_id, "")
        lead.refresh_from_db()
        self.assertEqual((lead.status, lead.converted_patient), (Lead.Status.BOOKED, patient))
        # A second expected patient without an ID is allowed too.
        other = Lead.objects.create(branch=self.branch, full_name="منى محمود علي", phone_primary="01005556668")
        self.client.post(f"/patients/calls/{other.pk}/book/")
        self.assertEqual(Patient.objects.filter(is_expected=True).count(), 2)
        # "Patient came": the short file is completed, without waiting for an approval.
        self.assertRedirects(self.client.get(f"/patients/new/?lead={lead.pk}"), f"/patients/{patient.pk}/edit/",
                             fetch_redirect_response=False)
        response = self.client.post(f"/patients/{patient.pk}/edit/", {
            "full_name": "سعيد محمود علي حسن", "id_type": "nid", "national_id": "29001150101234",
            "phone_primary": "01005556667", "phone_secondary": "01101234567", "preferred_phone": "primary",
            "marital_status": "married", "occupation": "مهندس", "city": "مدينة نصر", "missing_teeth": "single",
            "referral_source": self.facebook.pk, "status": "active"})
        self.assertRedirects(response, f"/patients/{patient.pk}/", fetch_redirect_response=False)
        patient.refresh_from_db()
        lead.refresh_from_db()
        self.assertFalse(patient.is_expected)
        self.assertEqual(patient.national_id, "29001150101234")
        self.assertEqual(lead.status, Lead.Status.CONVERTED)


class Round15MedicalRecallTests(TestCase):
    """Round 15: a high HbA1c gives a call for a new test after 3 months (the days are in Settings); the reception
    writes what the patient said and the next call date; a new reading closes it."""

    def setUp(self):
        from decimal import Decimal

        from apps.charting.models import Examination

        self.branch = setup_clinic()
        self.patient = make_patient(self.branch)
        make_user("sec", "secretary")
        self.old = timezone.localdate() - timedelta(days=100)
        self.exam = Examination.objects.create(patient=self.patient, history_only=True, medical_taken=True,
                                               hba1c=Decimal("8.4"), hba1c_date=self.old, exam_date=self.old)
        self.client.login(username="sec", password=PASSWORD)

    def test_a_high_hba1c_is_called_after_three_months(self):
        from apps.patients.models import MedicalRecall

        page = self.client.get("/patients/medical-follow-up/")
        recall = MedicalRecall.objects.get()
        self.assertEqual((recall.reason, recall.reading, recall.due_on),
                         ("hba1c", "8.4%", self.old + timedelta(days=90)))
        self.assertEqual(page.context["recalls"], [recall])
        self.assertIn("https://wa.me/20", page.context["recalls"][0].whatsapp)
        self.assertContains(self.client.get("/"), "8.4%")  # the reception's home page
        # Seen again: no second call for the same reading.
        self.client.get("/patients/medical-follow-up/")
        self.assertEqual(MedicalRecall.objects.count(), 1)

    def test_the_answers_set_the_next_call(self):
        from apps.patients.medical import sync_recalls
        from apps.patients.models import MedicalRecall

        sync_recalls(Patient.objects.all())
        recall = MedicalRecall.objects.get()
        url = f"/patients/medical-follow-up/calls/{recall.pk}/"
        self.client.post(url, {"outcome": "no_answer"})
        recall.refresh_from_db()
        self.assertEqual((recall.calls, recall.due_on), (1, timezone.localdate() + timedelta(days=2)))
        self.client.post(url, {"outcome": "will_do", "notes": "بعد العيد"})
        recall.refresh_from_db()
        self.assertEqual((recall.calls, recall.due_on, recall.notes, recall.status),
                         (2, timezone.localdate() + timedelta(days=14), "بعد العيد", "open"))
        self.assertEqual(self.client.get("/patients/medical-follow-up/").context["recalls"], [])  # not due yet
        self.client.post(url, {"outcome": "refused"})
        recall.refresh_from_db()
        self.assertEqual(recall.status, "stopped")

    def test_a_new_reading_closes_the_call(self):
        from decimal import Decimal

        from apps.charting.models import Examination
        from apps.patients.medical import sync_recalls
        from apps.patients.models import MedicalRecall

        sync_recalls(Patient.objects.all())
        Examination.objects.create(patient=self.patient, history_only=True, medical_taken=True,
                                   hba1c=Decimal("6.5"), hba1c_date=timezone.localdate())
        sync_recalls(Patient.objects.all())
        self.assertEqual(MedicalRecall.objects.get().status, "done")
