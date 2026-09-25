"""Read a stock list from Excel (.xlsx) or CSV.

Accepts the simple lists the academy already keeps (one row per item with a quantity
and a name, like "Dental_Material_and_Instrument.xlsx") and lists with a title row
naming the columns: name / item, quantity / qty, category, unit, min.
"""

import csv
import io
from decimal import Decimal, InvalidOperation

from django.db import transaction

from apps.core.utils import normalize_digits

from .models import StockCategory, StockItem, StockMovement
from .services import record_movement

HEADERS = {
    "name": ("name", "item", "items", "description", "الصنف", "الاسم"),
    "quantity": ("quantity", "qty", "count", "الكمية", "العدد"),
    "category": ("category", "type", "الفئة", "التصنيف"),
    "unit": ("unit", "الوحدة"),
    "min": ("min", "minimum", "reorder", "reorder level", "الحد الأدنى"),
}


def read_rows(upload):
    name = upload.name.lower()
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook

        workbook = load_workbook(upload, read_only=True, data_only=True)
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(values_only=True):
                yield list(row)
    else:
        text = upload.read().decode("utf-8-sig", errors="replace")
        yield from csv.reader(io.StringIO(text))


def _number(value):
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value))
    text = normalize_digits(str(value)).strip()
    try:
        return Decimal(text) if text else None
    except InvalidOperation:
        return None


def parse(rows):
    """Rows -> list of dicts {name, quantity, category, unit, min}."""
    columns = None
    items = []
    for row in rows:
        cells = [c for c in row]
        if not any(c not in (None, "") for c in cells):
            continue
        lowered = [str(c).strip().lower() if c is not None else "" for c in cells]
        if columns is None and not items:
            found = {key: lowered.index(title) for key, titles in HEADERS.items() for title in titles if title in lowered}
            if "name" in found:
                columns = found
                continue
        if columns:
            def cell(key):
                index = columns.get(key)
                return cells[index] if index is not None and index < len(cells) else None
            name = str(cell("name") or "").strip()
            entry = {"name": name, "quantity": _number(cell("quantity")), "category": str(cell("category") or "").strip(),
                     "unit": str(cell("unit") or "").strip(), "min": _number(cell("min"))}
        else:
            numbers = [_number(c) for c in cells]
            texts = [str(c).strip() for c, n in zip(cells, numbers) if c not in (None, "") and n is None]
            entry = {"name": texts[0] if texts else "", "quantity": next((n for n in numbers if n is not None), None),
                     "category": "", "unit": "", "min": None}
        if entry["name"]:
            entry["name"] = " ".join(entry["name"].split())[:150]
            items.append(entry)
    return items


def _category(name, default):
    if not name:
        return default
    for category in StockCategory.objects.all():
        if name.lower() in (category.name_ar.lower(), (category.name_en or "").lower()):
            return category
    return StockCategory.objects.create(name_ar=name[:120], name_en=name[:120])


@transaction.atomic
def import_items(entries, default_category, user, as_count=True, note="Imported list"):
    """Create missing items and bring quantities in. Returns (created, updated)."""
    created = updated = 0
    for entry in entries:
        category = _category(entry["category"], default_category)
        item = StockItem.objects.filter(name__iexact=entry["name"]).first()
        if item is None:
            item = StockItem.objects.create(name=entry["name"], category=category, unit=entry["unit"],
                                            min_quantity=entry["min"] or 0, created_by=user)
            created += 1
        else:
            updated += 1
            if entry["min"] is not None:
                StockItem.objects.filter(pk=item.pk).update(min_quantity=entry["min"])
        quantity = entry["quantity"]
        if quantity is None:
            continue
        if as_count:
            if quantity != item.quantity:
                record_movement(item, StockMovement.Kind.COUNT, quantity, user, notes=note)
        elif quantity > 0:
            record_movement(item, StockMovement.Kind.IN, quantity, user, notes=note)
    return created, updated
