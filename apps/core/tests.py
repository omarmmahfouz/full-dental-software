from datetime import date

from django.core.exceptions import ValidationError
from django.test import TestCase
from django.utils import translation

from apps.core.models import Notification
from apps.core.notify import notify_roles
from apps.core.roles import SUPERVISOR, has_role
from apps.core.testing import PASSWORD, make_user, setup_clinic
from apps.core.utils import normalize_digits, normalize_phone, parse_egyptian_national_id, validate_phone


class UtilsTests(TestCase):
    def test_arabic_digits_are_converted(self):
        self.assertEqual(normalize_digits("٠١٠٠١٢٣٤٥٦٧"), "01001234567")
        self.assertEqual(normalize_digits("۰۱۲"), "012")

    def test_phone_normalisation_makes_duplicates_comparable(self):
        expected = "01001234567"
        for typed in ["01001234567", "+20 100 123 4567", "0020-100-123-4567", "201001234567",
                      "1001234567", "٠١٠٠١٢٣٤٥٦٧", " 0100 123 4567 "]:
            with self.subTest(typed=typed):
                self.assertEqual(normalize_phone(typed), expected)
        self.assertEqual(normalize_phone("+966 50 123 4567"), "+966501234567")

    def test_validate_phone(self):
        self.assertEqual(validate_phone("01551234567"), "01551234567")
        self.assertEqual(validate_phone("+966501234567"), "+966501234567")
        with self.assertRaises(ValidationError):
            validate_phone("01301234567")  # not an Egyptian mobile prefix
        with self.assertRaises(ValidationError):
            validate_phone("0223456789")  # landline not allowed as primary
        self.assertEqual(validate_phone("0223456789", mobile_only=False), "0223456789")

    def test_national_id_parsing(self):
        data = parse_egyptian_national_id("29001150101234")  # 13th digit 3 -> male
        self.assertEqual(data["birth_date"], date(1990, 1, 15))
        self.assertEqual(data["gender"], "M")
        self.assertEqual(parse_egyptian_national_id("30105220101242")["gender"], "F")
        self.assertEqual(parse_egyptian_national_id("30105220101242")["birth_date"], date(2001, 5, 22))
        for bad in ["2900115010123", "19001150101234", "29013450101234", "abc"]:
            with self.subTest(bad=bad), self.assertRaises(ValidationError):
                parse_egyptian_national_id(bad)


class RolesAndNotificationTests(TestCase):
    def setUp(self):
        setup_clinic()

    def test_roles_and_notify(self):
        supervisor = make_user("sup", SUPERVISOR)
        other = make_user("sec", "secretary")
        self.assertTrue(has_role(supervisor, SUPERVISOR))
        self.assertFalse(has_role(other, SUPERVISOR))
        notify_roles((SUPERVISOR,), "Test %(x)s", params={"x": "1"})
        self.assertEqual(Notification.objects.get(recipient=supervisor).title, "Test 1")
        self.assertFalse(Notification.objects.filter(recipient=other).exists())

    def test_login_required_everywhere(self):
        response = self.client.get("/patients/")
        self.assertRedirects(response, "/login/?next=/patients/", fetch_redirect_response=False)

    def test_notification_open_marks_read(self):
        user = make_user("sec", "secretary")
        note = Notification.objects.create(recipient=user, title="x", url="/patients/")
        self.client.login(username="sec", password=PASSWORD)
        response = self.client.get(f"/notifications/{note.pk}/")
        self.assertRedirects(response, "/patients/", fetch_redirect_response=False)
        note.refresh_from_db()
        self.assertIsNotNone(note.read_at)

    def test_dashboard_renders_for_each_role(self):
        for index, role in enumerate(["secretary", "dentist", "supervisor", "owner"]):
            make_user(f"u{index}", role)
            self.client.login(username=f"u{index}", password=PASSWORD)
            self.assertEqual(self.client.get("/").status_code, 200)


class LanguageTests(TestCase):
    def setUp(self):
        setup_clinic()

    def language_of_page(self, username):
        self.client.logout()
        self.client.login(username=username, password=PASSWORD)
        return self.client.get("/").context["LANGUAGE_CODE"]

    def test_secretary_arabic_dentist_english(self):
        make_user("sec", "secretary")
        make_user("dentist", "dentist")
        make_user("owner", "owner")
        self.assertEqual(self.language_of_page("sec"), "ar")
        self.assertEqual(self.language_of_page("dentist"), "en")
        self.assertEqual(self.language_of_page("owner"), "en")

    def test_chosen_language_is_kept_per_user(self):
        make_user("sec", "secretary")
        make_user("dentist", "dentist")
        self.client.login(username="dentist", password=PASSWORD)
        response = self.client.post("/i18n/setlang/", {"language": "ar", "next": "/patients/"})
        self.assertRedirects(response, "/patients/", fetch_redirect_response=False)
        self.assertEqual(self.language_of_page("dentist"), "ar")
        self.assertEqual(self.language_of_page("sec"), "ar")
        self.client.login(username="sec", password=PASSWORD)
        self.client.post("/i18n/setlang/", {"language": "en", "next": "https://evil.example/"})
        self.assertEqual(self.language_of_page("sec"), "en")
        self.assertEqual(self.language_of_page("dentist"), "ar")  # one user's choice does not change another's

    def test_notifications_and_names_in_each_users_language(self):
        from django.utils.translation import gettext_lazy

        from apps.core.notify import notify_users
        from apps.scheduling.models import Room

        secretary = make_user("sec", "secretary")
        dentist = make_user("dentist", "dentist")
        notify_users([secretary, dentist], gettext_lazy("Lab request %(number)s needs your review"), params={"number": "L-1"})
        self.assertEqual(Notification.objects.get(recipient=dentist).title, "Lab request L-1 needs your review")
        self.assertNotIn("Lab request", Notification.objects.get(recipient=secretary).title)
        self.client.login(username="dentist", password=PASSWORD)
        page = self.client.get("/schedule/rooms/")
        self.assertIn("Room 1", [str(room) for room, _cells in page.context["rows"]])
        self.assertTrue(Room.objects.filter(name="غرفة 1", name_en="Room 1").exists())


class AccessAndSettingsTests(TestCase):
    def setUp(self):
        self.branch = setup_clinic()
        self.owner = make_user("owner", "owner")
        self.secretary = make_user("sec", "secretary")

    def test_read_only_person(self):
        from apps.core.models import UserProfile

        UserProfile.objects.update_or_create(user=self.secretary, defaults={"read_only": True})
        self.client.login(username="sec", password=PASSWORD)
        page = self.client.get("/patients/")
        self.assertEqual(page.status_code, 200)
        self.assertTrue(page.context["read_only_here"])
        self.assertEqual(self.client.post("/patients/calls/new/", {"full_name": "x"}).status_code, 403)
        # Changing the language is always allowed.
        self.assertEqual(self.client.post("/i18n/setlang/", {"language": "en"}).status_code, 302)

    def test_role_access_matrix(self):
        from apps.core.models import AreaAccess

        self.client.login(username="owner", password=PASSWORD)
        self.client.post("/settings/access/", {"secretary__purchases": "hidden", "secretary__complaints": "read"})
        self.assertEqual(AreaAccess.objects.get(role="secretary", area="purchases").level, "hidden")
        self.client.login(username="sec", password=PASSWORD)
        self.assertEqual(self.client.get("/purchases/").status_code, 403)
        self.assertIn("purchases", self.client.get("/").context["hidden_areas"])
        self.assertEqual(self.client.get("/complaints/").status_code, 200)
        self.assertEqual(self.client.post("/complaints/new/", {}).status_code, 403)
        self.client.login(username="owner", password=PASSWORD)  # the owner is never limited
        self.assertEqual(self.client.get("/purchases/").status_code, 200)

    def test_time_limits(self):
        from datetime import timedelta

        from django.utils import timezone

        from apps.core.models import UserProfile

        profile, _created = UserProfile.objects.get_or_create(user=self.secretary)
        self.client.login(username="sec", password=PASSWORD)
        profile.access_until = timezone.localdate() - timedelta(days=1)
        profile.save()
        response = self.client.get("/patients/")
        self.assertRedirects(response, "/login/", fetch_redirect_response=False)
        response = self.client.post("/login/", {"username": "sec", "password": PASSWORD})
        self.assertEqual(response.context["form"].non_field_errors().as_data()[0].code, "time")
        today = timezone.localdate().weekday()
        profile.access_until, profile.access_days = None, str((today + 1) % 7)
        profile.save()
        with translation.override("en"):
            self.assertIn("day", profile.access_problem())
        profile.access_days = ""
        profile.save()
        self.assertEqual(profile.access_problem(), "")

    def test_owner_manages_people_and_lists(self):
        from apps.core.models import ClinicSettings, UserProfile
        from apps.surgery.models import ImplantSystem

        self.client.login(username="owner", password=PASSWORD)
        self.client.post("/settings/users/new/", {
            "username": "newsec", "first_name": "Sara", "roles": ["secretary"], "is_active": "on",
            "new_password": "pass-12345", "read_only": "on", "access_days": ["5", "6"],
            "access_start": "09:00", "access_end": "17:00",
        })
        profile = UserProfile.objects.get(user__username="newsec")
        self.assertEqual((profile.read_only, profile.access_days), (True, "5,6"))
        self.assertTrue(profile.user.groups.filter(name="secretary").exists())
        self.client.post("/settings/lists/implant_systems/new/", {"company": "Zimmer", "line": "TSV", "is_active": "on"})
        self.assertTrue(ImplantSystem.objects.filter(company="Zimmer").exists())
        self.client.post("/settings/options/", {
            "o-day_start": "09:00", "o-day_end": "17:00", "o-surgery_days": ["3", "4"], "o-dicom_email": "ciapts@gmail.com",
            "o-late_threshold_minutes": 15, "o-default_appointment_minutes": 45, "o-complaint_follow_up_days": 3,
            "o-stock_expiry_days": 30, "o-reminder_days_before": 2, "o-whatsapp_country_code": "20",
            "b-name_ar": "أكاديمية القاهرة لزراعة الأسنان", "b-name_en": "Cairo Implant Academy", "b-phone": "0223456789",
            "b-address": "Cairo",
        })
        self.assertEqual(ClinicSettings.get().late_threshold_minutes, 15)
        make_user("head", "head_cia")
        self.client.login(username="head", password=PASSWORD)
        self.assertEqual(self.client.get("/settings/users/").status_code, 403)
        self.assertEqual(self.client.get("/settings/lists/implant_systems/").status_code, 200)


class DemoDataTests(TestCase):
    def test_files_the_demo_data_needs_are_in_the_project(self):
        from apps.core.management.commands.load_demo_data import STOCK_LIST

        self.assertTrue(STOCK_LIST.exists(), STOCK_LIST)

    def test_existing_practice_data_is_kept_and_missing_logins_are_named(self):
        import io

        from django.core.management import CommandError, call_command

        from apps.core.testing import make_patient

        make_patient(setup_clinic())
        make_user("owner", "owner")
        with self.assertRaises(CommandError):
            call_command("load_demo_data", password="x", stdout=io.StringIO())
        out = io.StringIO()
        call_command("load_demo_data", password="x", if_empty=True, stdout=out)
        self.assertIn("kept as it is", out.getvalue())
        self.assertIn("missing: headcia, teamhead", out.getvalue())
        self.assertNotIn("owner,", out.getvalue())


class WidgetTests(TestCase):
    def test_dates_are_dd_mm_yyyy_and_times_come_in_quarters(self):
        from datetime import time

        from django import forms as dj_forms

        from apps.core.forms import StyledForm

        class Form(StyledForm):
            day = dj_forms.DateField()
            at = dj_forms.DateTimeField()
            start = dj_forms.TimeField()

        form = Form({"day": "٠٥/٠٩/٢٠٢٦", "at_0": "05/09/2026", "at_1": "09:45", "start": "10:30"})
        self.assertTrue(form.is_valid(), form.errors)
        self.assertEqual(form.cleaned_data["day"], date(2026, 9, 5))
        self.assertEqual((form.cleaned_data["at"].day, form.cleaned_data["at"].month), (5, 9))
        self.assertEqual(form.cleaned_data["start"], time(10, 30))
        html = str(Form(initial={"day": date(2026, 9, 5)})["day"])
        self.assertIn('value="05/09/2026"', html)
        options = str(Form()["start"])
        self.assertIn('value="09:15"', options)
        self.assertNotIn('value="09:10"', options)


class PersonAccessTests(TestCase):
    def test_academy_for_one_secretary_only(self):
        from apps.core.models import AreaAccess, PersonAreaAccess

        setup_clinic()
        make_user("owner", "owner")
        first, second = make_user("sec1", "secretary"), make_user("sec2", "secretary")
        self.client.login(username="owner", password=PASSWORD)
        form = self.client.get(f"/settings/users/{second.pk}/").context["form"]
        self.assertIn("area__academy", form.fields)
        data = {"username": "sec2", "first_name": "Sara", "roles": ["secretary"], "is_active": "on",
                "area__academy": "hidden"}
        self.client.post(f"/settings/users/{second.pk}/", data)
        self.assertEqual(PersonAreaAccess.objects.get(user=second).level, "hidden")
        self.client.login(username="sec2", password=PASSWORD)
        self.assertEqual(self.client.get("/academy/candidates/").status_code, 403)
        self.assertNotIn("overdue_installments", self.client.get("/").context)
        self.client.login(username="sec1", password=PASSWORD)
        self.assertEqual(self.client.get("/academy/candidates/").status_code, 200)
        # A person rule opens again what the role matrix closed, but never more than the role allows.
        AreaAccess.objects.create(role="secretary", area="academy", level="hidden")
        PersonAreaAccess.objects.create(user=first, area="academy", level="full")
        self.assertEqual(self.client.get("/academy/candidates/").status_code, 200)
        PersonAreaAccess.objects.create(user=first, area="surgery", level="full")
        self.assertEqual(self.client.get("/surgery/").status_code, 403)


class ApprovalTests(TestCase):
    def setUp(self):
        from apps.core.testing import make_patient

        self.branch = setup_clinic()
        self.secretary = make_user("sec", "secretary")
        self.head = make_user("head", "head_cia")
        self.patient = make_patient(self.branch, phone="01001234567")

    def patient_data(self, **changes):
        from apps.patients.models import ReferralSource

        data = {"full_name": self.patient.full_name, "id_type": "nid", "national_id": self.patient.national_id,
                "phone_primary": self.patient.phone_primary, "preferred_phone": "primary", "missing_teeth": "unknown",
                "referral_source": ReferralSource.objects.filter(asks_for_patient=False).first().pk, "status": "active"}
        data.update(changes)
        return data

    def test_reception_edits_wait_for_the_head_of_cia(self):
        from apps.core.models import ChangeRequest
        from apps.patients.models import MedicalCondition

        diabetes = MedicalCondition.objects.first()
        self.client.login(username="sec", password=PASSWORD)
        self.client.post(f"/patients/{self.patient.pk}/edit/",
                         self.patient_data(phone_primary="01101234567", medical_conditions=[diabetes.pk]))
        self.patient.refresh_from_db()
        self.assertEqual(self.patient.phone_primary, "01001234567")  # nothing changes yet
        change = ChangeRequest.objects.get()
        self.assertLessEqual({"phone_primary", "medical_conditions"}, {c["field"] for c in change.changes})
        self.assertTrue(Notification.objects.filter(recipient=self.head).exists())
        self.assertContains(self.client.get(f"/patients/{self.patient.pk}/"), "01101234567")  # shown as waiting
        self.assertEqual(self.client.get("/approvals/").status_code, 403)

        self.client.login(username="head", password=PASSWORD)
        self.assertEqual(self.client.get("/").context["pending_approvals"], 1)
        self.client.post(f"/approvals/{change.pk}/", {"action": "approve", "note": "ok 100%"})
        self.patient.refresh_from_db()
        self.assertEqual(self.patient.phone_primary, "01101234567")
        self.assertEqual(list(self.patient.medical_conditions.all()), [diabetes])
        self.assertTrue(Notification.objects.filter(recipient=self.secretary).exists())
        # The head of CIA edits directly.
        self.client.post(f"/patients/{self.patient.pk}/edit/", self.patient_data(phone_primary="01201234567",
                                                                              medical_conditions=[diabetes.pk]))
        self.patient.refresh_from_db()
        self.assertEqual((self.patient.phone_primary, ChangeRequest.objects.count()), ("01201234567", 1))

    def test_forgotten_visit_times_are_corrected_after_approval(self):
        from datetime import timedelta

        from django.utils import timezone

        from apps.core.models import ChangeRequest
        from apps.scheduling.models import Appointment

        # Yesterday at 10:00, so the visit never crosses midnight whatever time the tests run.
        start = timezone.localtime().replace(hour=10, minute=0, second=0, microsecond=0) - timedelta(days=1)
        appointment = Appointment.objects.create(branch=self.branch, patient=self.patient, scheduled_at=start)
        appointment.mark_arrived(start)
        appointment.save()
        local = timezone.localtime(start)
        times = {"arrived_at_0": local.strftime("%d/%m/%Y"), "arrived_at_1": local.strftime("%H:%M"),
                 "entered_room_at_0": local.strftime("%d/%m/%Y"),
                 "entered_room_at_1": (local + timedelta(minutes=20)).strftime("%H:%M"),
                 "left_at_0": local.strftime("%d/%m/%Y"), "left_at_1": (local + timedelta(minutes=80)).strftime("%H:%M"),
                 "reason": "forgot to press Left"}
        self.client.login(username="sec", password=PASSWORD)
        self.client.post(f"/schedule/appointments/{appointment.pk}/times/", times)
        appointment.refresh_from_db()
        self.assertIsNone(appointment.left_at)
        change = ChangeRequest.objects.get(kind="visit_times")
        self.client.login(username="head", password=PASSWORD)
        self.client.post(f"/approvals/{change.pk}/", {"action": "approve"})
        appointment.refresh_from_db()
        self.assertEqual((appointment.status, appointment.chair_minutes), ("completed", 60))
        bad = dict(times, left_at_1=local.strftime("%H:%M"))  # left before entering
        self.client.post(f"/schedule/appointments/{appointment.pk}/times/", bad)
        self.assertEqual(ChangeRequest.objects.count(), 1)

    def test_rejected_changes_do_nothing(self):
        from apps.core.models import ChangeRequest

        self.client.login(username="sec", password=PASSWORD)
        self.client.post(f"/patients/{self.patient.pk}/edit/", self.patient_data(full_name="اسم آخر"))
        change = ChangeRequest.objects.get()
        self.client.login(username="head", password=PASSWORD)
        self.client.post(f"/approvals/{change.pk}/", {"action": "reject"})
        self.patient.refresh_from_db()
        change.refresh_from_db()
        self.assertEqual((change.status, self.patient.full_name == "اسم آخر"), ("rejected", False))


class ProblemReportTests(TestCase):
    def setUp(self):
        setup_clinic()
        self.owner = make_user("owner", "owner")
        self.secretary = make_user("sec", "secretary")

    def test_staff_report_a_problem_and_the_owner_answers(self):
        from apps.core.models import ProblemReport

        self.secretary.profile.read_only = True  # even read-only staff can report a problem
        self.secretary.profile.save()
        self.client.login(username="sec", password=PASSWORD)
        response = self.client.post("/problems/report/", {"description": "The print button does nothing",
                                                          "page": "/patients/1/"},
                                    headers={"X-Requested-With": "XMLHttpRequest"})
        self.assertTrue(response.json()["ok"])
        self.assertFalse(self.client.post("/problems/report/", {"description": ""},
                                          headers={"X-Requested-With": "XMLHttpRequest"}).json()["ok"])
        report = ProblemReport.objects.get()
        self.assertEqual((report.page, report.reported_by, report.kind), ("/patients/1/", self.secretary, "reported"))
        self.assertTrue(Notification.objects.filter(recipient=self.owner, url="/problems/").exists())
        self.assertEqual(self.client.get("/problems/").status_code, 403)  # only the owner and head of CIA read them
        self.client.login(username="owner", password=PASSWORD)
        self.assertContains(self.client.get("/problems/"), "The print button does nothing")
        self.client.post(f"/problems/{report.pk}/", {"status": "solved", "answer": "Fixed in the new version"})
        report.refresh_from_db()
        self.assertEqual((report.status, report.handled_by), ("solved", self.owner))
        self.assertTrue(Notification.objects.filter(recipient=self.secretary, message="Fixed in the new version").exists())
        export = self.client.get("/problems/export/")
        self.assertIn("The print button does nothing", export.content.decode("utf-8-sig"))

    def test_page_errors_are_recorded_once_and_counted(self):
        from django.http import Http404
        from django.test import RequestFactory

        from apps.core.middleware import ErrorRecorderMiddleware
        from apps.core.models import ProblemReport

        middleware = ErrorRecorderMiddleware(lambda request: None)
        request = RequestFactory().get("/patients/5/")
        request.user = self.secretary
        for _ in range(3):
            self.assertIsNone(middleware.process_exception(request, ValueError("bad value")))
        middleware.process_exception(request, Http404())  # a missing page is not a software problem
        report = ProblemReport.objects.get()
        self.assertEqual((report.kind, report.times, report.page), ("automatic", 3, "/patients/5/"))
        self.assertIn("ValueError: bad value", report.error)
        self.assertEqual(Notification.objects.filter(recipient=self.owner, level="danger").count(), 1)

    def test_pages_have_the_leave_warning_and_report_pop_ups(self):
        self.client.login(username="sec", password=PASSWORD)
        page = self.client.get("/")
        self.assertContains(page, 'id="leave-warning"')
        self.assertContains(page, 'id="report-problem"')


class BackupAndExportTests(TestCase):
    def setUp(self):
        import shutil
        import tempfile

        from django.test import override_settings

        folder = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, folder, ignore_errors=True)
        override = override_settings(MEDIA_ROOT=f"{folder}/media", BACKUP_DIR=f"{folder}/backups")
        override.enable()
        self.addCleanup(override.disable)
        from apps.core.testing import make_dentist, make_patient

        self.branch = setup_clinic()
        self.owner = make_user("owner", "owner")
        make_user("head", "head_cia")
        make_user("sec", "secretary")
        make_dentist("dentist", kind="candidate")
        self.patient = make_patient(self.branch, name="مريض النسخة")

    def test_full_backup_has_data_excel_csv_and_files(self):
        import io
        import zipfile

        from django.core.files.base import ContentFile
        from openpyxl import load_workbook

        from apps.core.backup import create_backup
        from apps.patients.models import PatientDocument

        PatientDocument.objects.create(patient=self.patient, kind="other",
                                       file=ContentFile(b"%PDF-1.4", name="scan.pdf"))
        with zipfile.ZipFile(create_backup()) as bundle:
            names = bundle.namelist()
            self.assertTrue({"database.json", "excel/all-data.xlsx", "README.txt", "csv/patients.patient.csv"} <= set(names))
            self.assertTrue(any(n.startswith("media/") and n.endswith(".pdf") for n in names))
            self.assertIn("مريض النسخة", bundle.read("csv/patients.patient.csv").decode("utf-8-sig"))
            self.assertNotIn("pbkdf2", bundle.read("csv/auth.user.csv").decode("utf-8-sig"))  # no passwords
            workbook = load_workbook(io.BytesIO(bundle.read("excel/all-data.xlsx")), read_only=True)
            self.assertIn("Patients", workbook.sheetnames)
            rows = list(workbook["Patients"].iter_rows(values_only=True))
            self.assertIn("full name (as on ID)", rows[0])
            self.assertIn("مريض النسخة", rows[1])

    def test_restore_puts_the_backup_back(self):
        from apps.core.backup import create_backup, list_backups, restore_backup
        from apps.core.testing import make_patient
        from apps.patients.models import Patient

        backup = create_backup()
        make_patient(self.branch, name="بعد النسخة", nid="29001011234568", phone="01001234568")
        restore_backup(backup)
        self.assertEqual(list(Patient.objects.values_list("full_name", flat=True)), ["مريض النسخة"])
        self.assertTrue(self.client.login(username="owner", password=PASSWORD))  # logins come back too
        self.assertEqual(len(list_backups()), 2)  # the state before restoring was saved first

    def test_backup_page_is_for_the_owner(self):
        self.client.login(username="head", password=PASSWORD)
        self.assertEqual(self.client.get("/settings/backup/").status_code, 403)
        self.client.login(username="owner", password=PASSWORD)
        self.assertEqual(self.client.get("/settings/backup/").status_code, 200)
        self.client.post("/settings/backup/")
        page = self.client.get("/settings/backup/")
        name = page.context["backups"][0]["name"]
        self.assertEqual(self.client.get(f"/settings/backup/{name}/").status_code, 200)
        self.assertEqual(self.client.get("/settings/backup/../../etc/").status_code, 404)
        excel = self.client.get("/settings/backup/excel/")
        self.assertTrue(b"".join(excel.streaming_content).startswith(b"PK"))

    def test_patient_file_as_word(self):
        self.client.login(username="dentist", password=PASSWORD)
        response = self.client.get(f"/patients/{self.patient.pk}/word/")
        self.assertEqual(response.status_code, 200)
        self.assertIn(".docx", response["Content-Disposition"])
        self.client.login(username="sec", password=PASSWORD)
        self.assertEqual(self.client.get(f"/patients/{self.patient.pk}/word/").status_code, 403)
