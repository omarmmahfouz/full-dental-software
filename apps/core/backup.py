"""Backups and exports: everything the system holds, in forms any other program can read.

A full backup is one ZIP file:

    database.json      every record (Django format: put back with ``manage.py restore_backup``)
    database.sqlite3   the database file itself (trial / single-PC installations)
    excel/all-data.xlsx one sheet per table, readable column names and values
    csv/*.csv          the same tables as CSV (UTF-8), for any database or program
    media/             the uploaded files: ID scans, photos, stickers, invoices…
    README.txt         what is in the file and how to use it

Backups are kept in the ``BACKUP_DIR`` folder (data/backups by default); the newest
``BACKUP_KEEP`` are kept. Make one from Settings → Backup and export, or with
``python manage.py backup`` (e.g. every night from the Windows Task Scheduler).
"""

import csv
import io
import json
import os
import sqlite3
import tempfile
import zipfile
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path

from django.apps import apps
from django.conf import settings
from django.core.management import call_command
from django.db import connection
from django.utils import timezone, translation

# Tables that Django rebuilds by itself, or that are only temporary.
SKIPPED = {"contenttypes.contenttype", "auth.permission", "sessions.session", "admin.logentry"}
SECRET_FIELDS = {"password"}  # kept in database.json (needed to log in again), never in Excel or CSV


def backup_dir():
    folder = Path(getattr(settings, "BACKUP_DIR", Path(settings.BASE_DIR) / "data" / "backups"))
    folder.mkdir(parents=True, exist_ok=True)
    return folder


def data_models():
    """The tables to export, the system's own first (patients, schedule…), then the logins."""
    own = [m for m in apps.get_models() if m.__module__.startswith("apps.")]
    others = [m for m in apps.get_models() if not m.__module__.startswith("apps.")
              and m._meta.label_lower not in SKIPPED]
    return own + others


def _columns(model):
    columns = []
    for field in model._meta.get_fields():
        if field.auto_created and not field.concrete:
            continue  # reverse relations
        if field.name in SECRET_FIELDS:
            continue
        if field.many_to_many:
            columns.append((field, str(field.verbose_name), "m2m"))
        elif field.is_relation:
            columns.append((field, f"{field.verbose_name} (id)", "id"))
            columns.append((field, str(field.verbose_name), "label"))
        elif field.concrete:
            columns.append((field, str(field.verbose_name), "value"))
    return columns


class _Labels:
    """str() of related records, read once per table (readable Excel without thousands of queries)."""

    def __init__(self):
        self.cache = {}

    def get(self, model, pk):
        if pk is None:
            return ""
        if model not in self.cache:
            self.cache[model] = {}
        labels = self.cache[model]
        if pk not in labels:
            if not labels:
                labels.update({obj.pk: str(obj) for obj in model._default_manager.all()[:20000]})
            if pk not in labels:
                obj = model._default_manager.filter(pk=pk).first()
                labels[pk] = str(obj) if obj else ""
        return labels[pk]


def _cell(value):
    if isinstance(value, datetime):
        return timezone.localtime(value).replace(tzinfo=None) if timezone.is_aware(value) else value
    if isinstance(value, (date, time, int, float, bool)) or value is None:
        return value
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    if hasattr(value, "name") and isinstance(getattr(value, "name", None), str):  # files: path inside media/
        return value.name
    return str(value)


def table_rows(model, labels):
    """(headers, rows) of one table, with readable values (choice labels, names of related records)."""
    columns = _columns(model)
    headers = [header for _field, header, _kind in columns]
    queryset = model._default_manager.all().order_by("pk")
    m2m = [field.name for field, _header, kind in columns if kind == "m2m"]
    if m2m:
        queryset = queryset.prefetch_related(*m2m)
    rows = []
    for obj in queryset.iterator(chunk_size=2000) if not m2m else queryset:
        row = []
        for field, _header, kind in columns:
            if kind == "m2m":
                row.append(", ".join(str(item) for item in getattr(obj, field.name).all()))
            elif kind == "id":
                row.append(getattr(obj, field.attname))
            elif kind == "label":
                row.append(labels.get(field.related_model, getattr(obj, field.attname)))
            elif field.choices:
                row.append(str(dict(field.flatchoices).get(getattr(obj, field.attname), getattr(obj, field.attname))))
            else:
                row.append(_cell(getattr(obj, field.attname)))
        rows.append(row)
    return headers, rows


def _sheet_name(model, used):
    name = "".join(c for c in str(model._meta.verbose_name_plural).capitalize() if c not in "[]:*?/\\")[:31]
    base, number = name, 2
    while name.lower() in used:
        suffix = f" {number}"
        name, number = base[:31 - len(suffix)] + suffix, number + 1
    used.add(name.lower())
    return name


def excel_workbook(output):
    """Write all tables to ``output`` (a path or file) as one Excel workbook: a sheet per table."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    workbook = Workbook(write_only=False)
    contents = workbook.active
    contents.title = "Contents"
    contents.append(["Table", "Sheet", "Records", "Exported"])
    contents["A1"].font = contents["B1"].font = contents["C1"].font = Font(bold=True)
    labels, used = _Labels(), {"contents"}
    with translation.override("en"):
        for model in data_models():
            headers, rows = table_rows(model, labels)
            sheet = workbook.create_sheet(_sheet_name(model, used))
            sheet.append(headers)
            for cell in sheet[1]:
                cell.font = Font(bold=True)
            for row in rows:
                sheet.append([_excel_safe(value) for value in row])
            sheet.freeze_panes = "A2"
            for index, header in enumerate(headers, start=1):
                sheet.column_dimensions[get_column_letter(index)].width = min(max(len(header) + 2, 10), 40)
            contents.append([model._meta.label, sheet.title, len(rows), timezone.localtime().strftime("%d/%m/%Y %H:%M")])
    contents.column_dimensions["A"].width = 34
    contents.column_dimensions["B"].width = 34
    workbook.save(output)


def _excel_safe(value):
    if isinstance(value, str):
        value = "".join(ch for ch in value if ch in "\t\n\r" or ord(ch) >= 32)  # Excel refuses control characters
        if value[:1] in ("=", "+", "-", "@") and len(value) > 1 and not value.lstrip("-+").replace(".", "").isdigit():
            return "'" + value  # never run as a formula
    return value


def write_csv_files(bundle, folder="csv"):
    labels = _Labels()
    with translation.override("en"):
        for model in data_models():
            headers, rows = table_rows(model, labels)
            text = io.StringIO()
            writer = csv.writer(text)
            writer.writerow(headers)
            for row in rows:
                writer.writerow(["" if value is None else value for value in row])
            bundle.writestr(f"{folder}/{model._meta.app_label}.{model._meta.model_name}.csv",
                            "﻿" + text.getvalue())


README = """CIA dental system — full backup made on {when}

database.json      Every record. To put it back into the system (same or newer version):
                     python manage.py restore_backup "{name}"
                   It is the standard Django fixture format: a developer can also load it with
                   "python manage.py loaddata database.json" or read it with any JSON tool.
database.sqlite3   The database file itself (only when the system uses SQLite, as in the trial).
excel/all-data.xlsx  One sheet per table with readable column names. "(id)" columns are the
                   numbers that link the tables (e.g. an appointment's patient (id) is the id
                   of a row in the Patients sheet).
csv/               The same tables as CSV files (UTF-8), to import into any database or program.
media/             All uploaded files: ID scans, clinical photos (media/Patient photos/…),
                   implant stickers, invoices, payment scans.

Keep a copy of this file outside the clinic PC (external disk, USB, or cloud drive).
"""


def create_backup(stdout=None):
    """Make a full backup ZIP in the backup folder; returns its path. Old backups beyond BACKUP_KEEP are removed."""
    folder = backup_dir()
    stamp = timezone.localtime().strftime("%Y-%m-%d_%H-%M-%S")
    path, number = folder / f"backup_{stamp}.zip", 1
    while path.exists():  # never overwrite a backup (e.g. the one being restored)
        number += 1
        path = folder / f"backup_{stamp}_{number}.zip"
    partial = path.with_suffix(".zip.part")
    with zipfile.ZipFile(partial, "w", zipfile.ZIP_DEFLATED, allowZip64=True) as bundle:
        dump = io.StringIO()
        call_command("dumpdata", exclude=sorted(SKIPPED), natural_foreign=True, indent=1, stdout=dump)
        bundle.writestr("database.json", dump.getvalue())
        if connection.vendor == "sqlite":
            copy = _sqlite_copy()
            try:
                bundle.write(copy, "database.sqlite3")
            finally:
                os.unlink(copy)
        with tempfile.TemporaryFile() as workbook:
            excel_workbook(workbook)
            workbook.seek(0)
            bundle.writestr("excel/all-data.xlsx", workbook.read())
        write_csv_files(bundle)
        media_root = Path(settings.MEDIA_ROOT)
        if media_root.exists():
            for file in sorted(media_root.rglob("*")):
                if file.is_file() and folder not in file.parents:
                    bundle.write(file, f"media/{file.relative_to(media_root).as_posix()}", zipfile.ZIP_STORED)
        bundle.writestr("README.txt", README.format(when=timezone.localtime().strftime("%d/%m/%Y %H:%M"),
                                                    name=path.name))
    os.replace(partial, path)
    keep = getattr(settings, "BACKUP_KEEP", 10)
    for old in list_backups()[keep:]:
        old["path"].unlink(missing_ok=True)
    if stdout is not None:
        stdout.write(f"Backup saved: {path}")
    return path


def _sqlite_copy():
    """A consistent copy of the SQLite database, taken through the system's own connection
    (safe while the system is in use)."""
    handle = tempfile.NamedTemporaryFile(suffix=".sqlite3", delete=False)
    handle.close()
    connection.ensure_connection()
    target = sqlite3.connect(handle.name)
    try:
        if connection.in_atomic_block:  # inside a transaction SQLite's backup would wait for it to end
            target.executescript("\n".join(connection.connection.iterdump()))
        else:
            connection.connection.backup(target)
    finally:
        target.close()
    return handle.name


def list_backups():
    files = sorted(backup_dir().glob("backup_*.zip"), key=lambda p: p.stat().st_mtime, reverse=True)
    return [{"path": p, "name": p.name, "size_mb": round(p.stat().st_size / 1024 / 1024, 1),
             "made_at": timezone.make_aware(datetime.fromtimestamp(p.stat().st_mtime))} for p in files]


def restore_backup(path, stdout=None):
    """Replace all data and uploaded files with those of a backup ZIP (a backup of the current
    state is made first). Used after moving to a new PC or after a big change went wrong."""
    safety = create_backup()
    with zipfile.ZipFile(path) as bundle:
        with tempfile.TemporaryDirectory() as work:
            fixture = Path(work) / "database.json"
            fixture.write_bytes(bundle.read("database.json"))
            call_command("flush", interactive=False, verbosity=0)
            call_command("loaddata", str(fixture), verbosity=0)
            media_root = Path(settings.MEDIA_ROOT)
            for member in bundle.namelist():
                if member.startswith("media/") and not member.endswith("/"):
                    target = (media_root / member[len("media/"):]).resolve()
                    if media_root.resolve() not in target.parents:
                        continue  # never write outside the media folder
                    target.parent.mkdir(parents=True, exist_ok=True)
                    target.write_bytes(bundle.read(member))
    if stdout is not None:
        stdout.write(f"Restored from {path}. The data before restoring was saved in {safety}.")
    return safety

